#!/usr/bin/env python3
"""Generate an ATS-optimized one-page resume (DOCX + PDF) for Amit Kumar Mohanty,
tailored to the HCL 'AI Developer I' (graduate GenAI engineer) role."""

import os

# ---------------------------------------------------------------------------
# CONTENT (single source of truth for both DOCX and PDF)
# ---------------------------------------------------------------------------
NAME = "Amit Kumar Mohanty"
CONTACT = "mohantyamit2003@gmail.com | +91 9354937256 | github.com/amitmohanty022 | linkedin.com/in/amitkrmohanty"
TARGET = "AI Developer I  |  HCLTech"

SUMMARY = (
    "AI/ML Engineer and Computer Science graduate (AI/ML specialization) with 1+ year of hands-on "
    "experience building Generative AI (GenAI) applications and deploying models in cloud-native "
    "environments. Skilled in Python, GenAI frameworks and LLM APIs (OpenAI, Gemini, LangChain), "
    "prompt engineering, and containerized microservices (Docker, Kubernetes, CI/CD). Proven ability "
    "to take solutions from ideation to MVP within Agile/Scrum teams. Solution-oriented, strong "
    "communicator, and eager to pursue AI/ML and cloud certifications from AWS, Google, Microsoft, and Red Hat."
)

EDUCATION = [
    ("Birla Institute of Technology and Science, Pilani (BITS Pilani)", "July 2026 - July 2028",
     "M.Tech in Artificial Intelligence and Machine Learning (Pursuing)"),
    ("K.R. Mangalam University, Gurgaon", "July 2021 - July 2025",
     "B.Tech in Computer Science and Engineering (Specialization in AI/ML)"),
]

EXPERIENCE = [
    ("Research Associate - AI, Keywords Studios India", "Gurgaon, Haryana", "May 2025 - Present", [
        "Trained and optimized enterprise-grade Generative AI and agentic AI models to autonomously execute complex multi-step workflows, improving task-completion reliability across cloud-based deployments.",
        "Contributed to a web-automation AI agent that interprets on-screen content and, from natural-language prompts, plans and executes browser actions to complete multi-step tasks.",
        "Built custom data pipelines and delivered 85% of self-generated training datasets to production, applying data preprocessing and feature engineering for high-quality model inputs.",
        "Awarded \"Rookie of the Year 2025\" for surpassing all KPIs through model training, data-pipeline optimization, and cross-functional collaboration in an Agile team.",
    ]),
    ("Artificial Intelligence Intern, Infosys", "Remote", "Nov 2024 - Feb 2025", [
        "Engineered an AI call-center system with LangChain and LangSmith, boosting model response accuracy by 35% and improving customer-AI interaction quality.",
        "Designed and A/B-tested prompt-engineering strategies with automated evaluation, reducing hallucinations by 20% and mitigating biased/undesired model outputs.",
        "Developed and deployed LLM features that reduced system latency by 25%, streamlining the model deployment process within a 12-engineer team.",
    ]),
]

PROJECTS = [
    ("Dynamic Screen Companion", "June 2025 - Present", [
        "Built a real-time multi-modal GenAI assistant integrating the Gemini API to analyze live screen activity and deliver continuous voice feedback.",
        "Engineered a high-performance FastAPI microservice backend handling concurrent inference streams, containerized with Docker and CI/CD for cloud-native deployment.",
        "Optimized OCR and NLP pipelines (tokenization, embeddings) to maintain 95%+ extraction accuracy across varied screen layouts.",
    ]),
    ("Diabetic Optiscan (Peer-reviewed research)", "Jan 2025 - June 2025", [
        "Built an AI tool detecting early diabetic retinopathy from retinal images; co-authored a peer-reviewed paper on the design and accuracy gains.",
        "Led integration of a Vision Transformer (ViT-B/16) on the APTOS dataset, optimizing self-attention to capture complex spatial patterns.",
        "Designed a custom wavelet technique that boosted detection precision by 50% over legacy methods.",
    ]),
]

SKILLS = [
    ("Languages", "Python, Java (basic), C++, SQL, R"),
    ("Generative AI & NLP", "LLMs, GenAI, OpenAI API, Gemini API, LangChain, LangGraph, RAG, Prompt Engineering, Fine-tuning (LoRA/QLoRA), Transformers, Tokenization, Embeddings, Hugging Face, BERT, AI Ethics & Bias Mitigation"),
    ("Cloud-Native & DevOps", "Docker, Kubernetes, Microservices, CI/CD, 12-Factor Apps, MLflow, Model Serving & Monitoring, Git, Agile/Scrum"),
    ("Data & Cloud", "AWS, GCP (Vertex AI), Data Pipelines, Feature Engineering, Vector Databases (FAISS/Chroma), Pandas, NumPy, MySQL, MongoDB, Power BI"),
    ("Frameworks", "FastAPI, TensorFlow, PyTorch, Scikit-learn, OpenCV"),
]

CERTIFICATIONS = [
    "Ducat India - Data Science Professional Training (Python, ML, statistics, predictive modeling), Apr 2024 - Apr 2025.",
    "ISRO - AI/ML for Geodata Analysis (applied ML/DL to remote-sensing imagery).",
    "Actively pursuing cloud & AI/ML certifications (AWS, Google Cloud, Microsoft Azure, Red Hat).",
]

# ---------------------------------------------------------------------------
# DOCX GENERATION
# ---------------------------------------------------------------------------
def build_docx(path):
    from docx import Document
    from docx.shared import Pt, RGBColor, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    doc = Document()
    # Margins
    for s in doc.sections:
        s.top_margin = Inches(0.4)
        s.bottom_margin = Inches(0.4)
        s.left_margin = Inches(0.6)
        s.right_margin = Inches(0.6)

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(9.5)
    style.paragraph_format.space_after = Pt(0)
    style.paragraph_format.space_before = Pt(0)
    style.paragraph_format.line_spacing = 1.0

    def set_spacing(p, before=0, after=0, line=1.0):
        pf = p.paragraph_format
        pf.space_before = Pt(before)
        pf.space_after = Pt(after)
        pf.line_spacing = line

    # Name
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(NAME)
    r.bold = True
    r.font.size = Pt(20)
    set_spacing(p, 0, 1)

    # Contact
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(CONTACT)
    r.font.size = Pt(9)
    set_spacing(p, 0, 1)

    # Target
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Target Role: " + TARGET)
    r.font.size = Pt(9.5)
    r.italic = True
    set_spacing(p, 0, 4)

    def heading(text):
        p = doc.add_paragraph()
        r = p.add_run(text.upper())
        r.bold = True
        r.font.size = Pt(11)
        r.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
        set_spacing(p, 5, 1)
        # bottom border
        pPr = p._p.get_or_add_pPr()
        pbdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "6")
        bottom.set(qn("w:space"), "1")
        bottom.set(qn("w:color"), "1F3A5F")
        pbdr.append(bottom)
        pPr.append(pbdr)

    def bullet(text):
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Inches(0.18)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(text)
        r.font.size = Pt(9.5)
        return p

    def title_line(left_bold, right):
        p = doc.add_paragraph()
        set_spacing(p, 2, 0)
        # tab stop at right margin
        from docx.enum.text import WD_TAB_ALIGNMENT
        section_w = Inches(8.5) - Inches(0.6) - Inches(0.6)
        p.paragraph_format.tab_stops.add_tab_stop(section_w, WD_TAB_ALIGNMENT.RIGHT)
        r = p.add_run(left_bold)
        r.bold = True
        r.font.size = Pt(10)
        r2 = p.add_run("\t" + right)
        r2.font.size = Pt(9)
        r2.bold = True
        return p

    def sub_line(text):
        p = doc.add_paragraph()
        set_spacing(p, 0, 0)
        r = p.add_run(text)
        r.italic = True
        r.font.size = Pt(9.5)

    # SUMMARY
    heading("Professional Summary")
    p = doc.add_paragraph()
    set_spacing(p, 2, 0)
    r = p.add_run(SUMMARY)
    r.font.size = Pt(9.5)

    # EDUCATION
    heading("Education")
    for school, dates, degree in EDUCATION:
        title_line(school, dates)
        sub_line(degree)

    # EXPERIENCE
    heading("Experience")
    for role, loc, dates, bullets in EXPERIENCE:
        title_line(role + "  -  " + loc, dates)
        for b in bullets:
            bullet(b)

    # PROJECTS
    heading("Projects")
    for name, dates, bullets in PROJECTS:
        title_line(name, dates)
        for b in bullets:
            bullet(b)

    # SKILLS
    heading("Skills")
    for cat, items in SKILLS:
        p = doc.add_paragraph()
        set_spacing(p, 1, 0)
        r = p.add_run(cat + ": ")
        r.bold = True
        r.font.size = Pt(9.5)
        r2 = p.add_run(items)
        r2.font.size = Pt(9.5)

    # CERTIFICATIONS
    heading("Certifications")
    for c in CERTIFICATIONS:
        bullet(c)

    doc.save(path)


# ---------------------------------------------------------------------------
# PDF GENERATION
# ---------------------------------------------------------------------------
def build_pdf(path):
    from fpdf import FPDF

    NAVY = (31, 58, 95)
    BLACK = (0, 0, 0)

    pdf = FPDF(orientation="P", unit="mm", format="A4")
    pdf.set_auto_page_break(auto=False)
    pdf.set_margins(left=15, top=10, right=15)
    pdf.add_page()
    epw = pdf.w - 30  # effective page width

    def name_header():
        pdf.set_font("Helvetica", "B", 22)
        pdf.set_text_color(*BLACK)
        pdf.cell(0, 9, NAME, align="C", new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("Helvetica", "", 8.5)
        pdf.cell(0, 4.5, CONTACT, align="C", new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("Helvetica", "I", 9)
        pdf.set_text_color(*NAVY)
        pdf.cell(0, 4.5, "Target Role: " + TARGET, align="C", new_x="LMARGIN", new_y="NEXT")
        pdf.ln(1.5)

    def heading(text):
        pdf.ln(1.2)
        pdf.set_font("Helvetica", "B", 10.5)
        pdf.set_text_color(*NAVY)
        y = pdf.get_y()
        pdf.cell(0, 5, text.upper(), new_x="LMARGIN", new_y="NEXT")
        pdf.set_draw_color(*NAVY)
        pdf.set_line_width(0.3)
        pdf.line(15, pdf.get_y(), 15 + epw, pdf.get_y())
        pdf.ln(0.8)
        pdf.set_text_color(*BLACK)

    def body(text, size=9.3):
        pdf.set_font("Helvetica", "", size)
        pdf.set_text_color(*BLACK)
        pdf.multi_cell(0, 4.2, text, new_x="LMARGIN", new_y="NEXT")

    def title_line(left, right):
        pdf.set_font("Helvetica", "B", 10)
        pdf.set_text_color(*BLACK)
        rw = pdf.get_string_width(right) + 1
        pdf.cell(epw - rw, 4.8, left, new_x="RIGHT", new_y="TOP")
        pdf.set_font("Helvetica", "B", 8.8)
        pdf.cell(rw, 4.8, right, align="R", new_x="LMARGIN", new_y="NEXT")

    def sub_line(text):
        pdf.set_font("Helvetica", "I", 9.2)
        pdf.multi_cell(0, 4.2, text, new_x="LMARGIN", new_y="NEXT")

    def bullet(text, size=9.3):
        pdf.set_font("Helvetica", "", size)
        pdf.set_text_color(*BLACK)
        x0 = pdf.get_x()
        pdf.cell(4, 4.2, chr(149), new_x="RIGHT", new_y="TOP")
        pdf.set_x(x0 + 4)
        pdf.multi_cell(epw - 4, 4.2, text, new_x="LMARGIN", new_y="NEXT")

    def cat_line(cat, items, size=9.3):
        pdf.set_x(15)
        pdf.set_font("Helvetica", "B", size)
        cw = pdf.get_string_width(cat + ": ")
        pdf.cell(cw, 4.2, cat + ": ", new_x="RIGHT", new_y="TOP")
        pdf.set_x(15 + cw)
        pdf.set_font("Helvetica", "", size)
        pdf.multi_cell(epw - cw, 4.2, items, new_x="LMARGIN", new_y="NEXT")

    name_header()

    heading("Professional Summary")
    body(SUMMARY)

    heading("Education")
    for school, dates, degree in EDUCATION:
        title_line(school, dates)
        sub_line(degree)

    heading("Experience")
    for role, loc, dates, bullets in EXPERIENCE:
        title_line(role + "  -  " + loc, dates)
        for b in bullets:
            bullet(b)
        pdf.ln(0.6)

    heading("Projects")
    for name, dates, bullets in PROJECTS:
        title_line(name, dates)
        for b in bullets:
            bullet(b)
        pdf.ln(0.6)

    heading("Skills")
    for cat, items in SKILLS:
        cat_line(cat, items)

    heading("Certifications")
    for c in CERTIFICATIONS:
        bullet(c)

    pdf.output(path)
    return pdf.page_no()


if __name__ == "__main__":
    from resume_lib import build_docx as _bd, build_pdf as _bp, autofit_opts
    CONTENT = {
        "NAME": NAME, "CONTACT": CONTACT, "TARGET": TARGET, "SUMMARY": SUMMARY,
        "EDUCATION": EDUCATION, "EXPERIENCE": EXPERIENCE, "PROJECTS": PROJECTS,
        "SKILLS": SKILLS, "CERTIFICATIONS": CERTIFICATIONS,
    }
    out_dir = os.path.dirname(os.path.abspath(__file__))
    docx_path = os.path.join(out_dir, "Amit Kumar Mohanty - HCLTech AI Developer I.docx")
    pdf_path = os.path.join(out_dir, "Amit Kumar Mohanty - HCLTech AI Developer I.pdf")
    opts = autofit_opts(CONTENT)
    _bd(CONTENT, docx_path, opts)
    pages, last_y = _bp(CONTENT, pdf_path, opts)
    print("DOCX ->", docx_path)
    print("PDF  ->", pdf_path, "| pages:", pages, "| last_y(mm):", round(last_y, 1))
