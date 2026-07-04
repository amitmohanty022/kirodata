#!/usr/bin/env python3
"""Reusable ATS-friendly one-page resume renderer (DOCX + PDF).

A `content` dict drives both renderers. Expected keys:
    NAME, CONTACT, TARGET, SUMMARY (str)
    EDUCATION: list of (school, dates, degree)
    EXPERIENCE: list of (role, loc, dates, [bullets])
    PROJECTS:   list of (name, dates, [bullets])
    SKILLS:     list of (category, items)
    CERTIFICATIONS: list of str

An optional `opts` dict tunes sizing/spacing so the content fills a single page
without leaving large whitespace at the bottom (while staying ATS-clean).
"""

DEFAULT_OPTS = {
    "name_size": 22.0,
    "contact_size": 8.5,
    "target_size": 9.0,
    "head_size": 10.5,
    "body_size": 9.3,
    "title_size": 10.0,
    "date_size": 8.8,
    "sub_size": 9.2,
    "lh": 4.2,          # line height (mm) for body / bullets / sub
    "title_h": 4.8,     # line height (mm) for title lines
    "sec_before": 1.2,  # blank space (mm) before each heading
    "sec_after": 0.8,   # blank space (mm) after heading rule
    "entry_gap": 0.6,   # blank space (mm) after each experience/project entry
    "header_gap": 1.5,  # blank space (mm) after the header block
    # docx-specific paragraph spacing (points)
    "docx_body_size": 9.5,
    "docx_bullet_after": 0.0,
    "docx_entry_after": 0.0,
    "docx_line": 1.0,
}


def _merge(opts):
    o = dict(DEFAULT_OPTS)
    if opts:
        o.update(opts)
    return o


# ---------------------------------------------------------------------------
# DOCX GENERATION
# ---------------------------------------------------------------------------
def build_docx(c, path, opts=None):
    from docx import Document
    from docx.shared import Pt, RGBColor, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    o = _merge(opts)
    bsz = o["docx_body_size"]

    doc = Document()
    for s in doc.sections:
        s.top_margin = Inches(0.5)
        s.bottom_margin = Inches(0.5)
        s.left_margin = Inches(0.6)
        s.right_margin = Inches(0.6)

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(bsz)
    style.paragraph_format.space_after = Pt(0)
    style.paragraph_format.space_before = Pt(0)
    style.paragraph_format.line_spacing = o["docx_line"]

    def set_spacing(p, before=0, after=0, line=None):
        pf = p.paragraph_format
        pf.space_before = Pt(before)
        pf.space_after = Pt(after)
        pf.line_spacing = line if line is not None else o["docx_line"]

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(c["NAME"]); r.bold = True; r.font.size = Pt(o["name_size"])
    set_spacing(p, 0, 2)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(c["CONTACT"]); r.font.size = Pt(bsz - 0.5)
    set_spacing(p, 0, 2)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Target Role: " + c["TARGET"]); r.font.size = Pt(bsz); r.italic = True
    set_spacing(p, 0, 5)

    def heading(text):
        p = doc.add_paragraph()
        r = p.add_run(text.upper()); r.bold = True; r.font.size = Pt(o["head_size"] + 0.5)
        r.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
        set_spacing(p, 6, 2)
        pPr = p._p.get_or_add_pPr()
        pbdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single"); bottom.set(qn("w:sz"), "6")
        bottom.set(qn("w:space"), "1"); bottom.set(qn("w:color"), "1F3A5F")
        pbdr.append(bottom); pPr.append(pbdr)

    def bullet(text):
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Inches(0.18)
        p.paragraph_format.space_after = Pt(o["docx_bullet_after"])
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.line_spacing = o["docx_line"]
        r = p.add_run(text); r.font.size = Pt(bsz)

    def title_line(left_bold, right):
        p = doc.add_paragraph()
        set_spacing(p, 3, 0)
        section_w = Inches(8.5) - Inches(0.6) - Inches(0.6)
        p.paragraph_format.tab_stops.add_tab_stop(section_w, WD_TAB_ALIGNMENT.RIGHT)
        r = p.add_run(left_bold); r.bold = True; r.font.size = Pt(o["title_size"])
        r2 = p.add_run("\t" + right); r2.font.size = Pt(o["date_size"] + 0.2); r2.bold = True

    def sub_line(text):
        p = doc.add_paragraph(); set_spacing(p, 0, 0)
        r = p.add_run(text); r.italic = True; r.font.size = Pt(bsz)

    def entry_gap():
        if o["docx_entry_after"] > 0:
            p = doc.add_paragraph(); set_spacing(p, 0, 0)
            r = p.add_run(""); r.font.size = Pt(o["docx_entry_after"])

    heading("Professional Summary")
    p = doc.add_paragraph(); set_spacing(p, 2, 0)
    r = p.add_run(c["SUMMARY"]); r.font.size = Pt(bsz)

    heading("Education")
    for school, dates, degree in c["EDUCATION"]:
        title_line(school, dates); sub_line(degree)

    heading("Experience")
    for role, loc, dates, bullets in c["EXPERIENCE"]:
        title_line(role + "  -  " + loc, dates)
        for b in bullets:
            bullet(b)
        entry_gap()

    heading("Projects")
    for name, dates, bullets in c["PROJECTS"]:
        title_line(name, dates)
        for b in bullets:
            bullet(b)
        entry_gap()

    heading("Skills")
    for cat, items in c["SKILLS"]:
        p = doc.add_paragraph(); set_spacing(p, 1, 0)
        r = p.add_run(cat + ": "); r.bold = True; r.font.size = Pt(bsz)
        r2 = p.add_run(items); r2.font.size = Pt(bsz)

    heading("Certifications")
    for cert in c["CERTIFICATIONS"]:
        bullet(cert)

    doc.save(path)


# ---------------------------------------------------------------------------
# PDF GENERATION
# ---------------------------------------------------------------------------
def build_pdf(c, path, opts=None):
    from fpdf import FPDF

    o = _merge(opts)
    NAVY = (31, 58, 95)
    BLACK = (0, 0, 0)
    lh = o["lh"]

    pdf = FPDF(orientation="P", unit="mm", format="A4")
    pdf.set_auto_page_break(auto=False)
    pdf.set_margins(left=15, top=11, right=15)
    pdf.add_page()
    epw = pdf.w - 30

    def name_header():
        pdf.set_font("Helvetica", "B", o["name_size"]); pdf.set_text_color(*BLACK)
        pdf.cell(0, o["name_size"] * 0.42, c["NAME"], align="C", new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("Helvetica", "", o["contact_size"])
        pdf.cell(0, 4.6, c["CONTACT"], align="C", new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("Helvetica", "I", o["target_size"]); pdf.set_text_color(*NAVY)
        pdf.cell(0, 4.6, "Target Role: " + c["TARGET"], align="C", new_x="LMARGIN", new_y="NEXT")
        pdf.ln(o["header_gap"])

    def heading(text):
        pdf.ln(o["sec_before"])
        pdf.set_font("Helvetica", "B", o["head_size"]); pdf.set_text_color(*NAVY)
        pdf.cell(0, 5.2, text.upper(), new_x="LMARGIN", new_y="NEXT")
        pdf.set_draw_color(*NAVY); pdf.set_line_width(0.3)
        pdf.line(15, pdf.get_y(), 15 + epw, pdf.get_y())
        pdf.ln(o["sec_after"]); pdf.set_text_color(*BLACK)

    def body(text):
        pdf.set_font("Helvetica", "", o["body_size"]); pdf.set_text_color(*BLACK)
        pdf.multi_cell(0, lh, text, new_x="LMARGIN", new_y="NEXT")

    def title_line(left, right):
        pdf.set_font("Helvetica", "B", o["title_size"]); pdf.set_text_color(*BLACK)
        rw = pdf.get_string_width(right) + 1
        pdf.cell(epw - rw, o["title_h"], left, new_x="RIGHT", new_y="TOP")
        pdf.set_font("Helvetica", "B", o["date_size"])
        pdf.cell(rw, o["title_h"], right, align="R", new_x="LMARGIN", new_y="NEXT")

    def sub_line(text):
        pdf.set_font("Helvetica", "I", o["sub_size"])
        pdf.multi_cell(0, lh, text, new_x="LMARGIN", new_y="NEXT")

    def bullet(text):
        pdf.set_font("Helvetica", "", o["body_size"]); pdf.set_text_color(*BLACK)
        x0 = pdf.get_x()
        pdf.cell(4, lh, chr(149), new_x="RIGHT", new_y="TOP")
        pdf.set_x(x0 + 4)
        pdf.multi_cell(epw - 4, lh, text, new_x="LMARGIN", new_y="NEXT")

    def cat_line(cat, items):
        pdf.set_x(15); pdf.set_font("Helvetica", "B", o["body_size"])
        cw = pdf.get_string_width(cat + ": ")
        pdf.cell(cw, lh, cat + ": ", new_x="RIGHT", new_y="TOP")
        pdf.set_x(15 + cw); pdf.set_font("Helvetica", "", o["body_size"])
        pdf.multi_cell(epw - cw, lh, items, new_x="LMARGIN", new_y="NEXT")

    name_header()
    heading("Professional Summary"); body(c["SUMMARY"])
    heading("Education")
    for school, dates, degree in c["EDUCATION"]:
        title_line(school, dates); sub_line(degree)
    heading("Experience")
    for role, loc, dates, bullets in c["EXPERIENCE"]:
        title_line(role + "  -  " + loc, dates)
        for b in bullets:
            bullet(b)
        pdf.ln(o["entry_gap"])
    heading("Projects")
    for name, dates, bullets in c["PROJECTS"]:
        title_line(name, dates)
        for b in bullets:
            bullet(b)
        pdf.ln(o["entry_gap"])
    heading("Skills")
    for cat, items in c["SKILLS"]:
        cat_line(cat, items)
    heading("Certifications")
    for cert in c["CERTIFICATIONS"]:
        bullet(cert)

    pdf.output(path)
    return pdf.page_no(), pdf.get_y()



# ---------------------------------------------------------------------------
# AUTO-FIT: render once, measure, then distribute spacing to fill the page
# ---------------------------------------------------------------------------
def _opts_for_scale(s, fixed_gaps):
    """Build an opts dict where fonts + line-heights are scaled by `s`.
    Gaps are held at fixed (small) values during the search phase."""
    o = {
        "name_size": 22.0 * s,
        "contact_size": 8.7 * s,
        "target_size": 9.2 * s,
        "head_size": 10.7 * s,
        "body_size": 9.7 * s,
        "title_size": 10.2 * s,
        "date_size": 8.9 * s,
        "sub_size": 9.5 * s,
        "lh": 4.55 * s,
        "title_h": 4.95 * s,
        "sec_before": fixed_gaps["sec_before"],
        "sec_after": fixed_gaps["sec_after"],
        "entry_gap": fixed_gaps["entry_gap"],
        "header_gap": fixed_gaps["header_gap"],
    }
    # DOCX mirrors the same visual density
    o["docx_body_size"] = max(9.5, min(11.5, 9.7 * s))
    o["docx_line"] = 1.0 + max(0.0, (s - 1.0)) * 0.15
    o["docx_bullet_after"] = 1.2 * s
    o["docx_entry_after"] = 3.0 * s
    return o


def _measure(content, opts):
    import os
    import tempfile
    fd, tmp = tempfile.mkstemp(suffix=".pdf")
    os.close(fd)
    _, y = build_pdf(content, tmp, opts)
    try:
        os.remove(tmp)
    except OSError:
        pass
    return y


def autofit_opts(content, target=286.0):
    """Return opts tuned so the PDF fills close to a full single page.

    Binary-searches a single scale factor over fonts + line height so the
    content naturally reaches the target bottom. Longer resumes shrink,
    shorter ones grow (within a readable range). If content is still too
    short at the max readable size, the remaining space is distributed as
    section/entry whitespace. All changes are additive whitespace / font
    size only, so ATS text extraction is never affected.
    """
    fixed = {"sec_before": 1.4, "sec_after": 0.9, "entry_gap": 0.6, "header_gap": 1.8}

    lo, hi = 0.70, 1.16
    # height is monotonic increasing in s
    best = None
    for _ in range(28):
        mid = (lo + hi) / 2.0
        y = _measure(content, _opts_for_scale(mid, fixed))
        best = (mid, y)
        if y > target:
            hi = mid
        else:
            lo = mid
    s = best[0]

    opts = _opts_for_scale(s, dict(fixed))
    y = _measure(content, opts)

    # If we hit the max scale but content is still short, fill with whitespace.
    extra = target - y
    if extra > 4.0 and s >= hi - 1e-3:
        n_head = 6
        n_entry = len(content["EXPERIENCE"]) + len(content["PROJECTS"])
        opts["sec_before"] += (extra * 0.55) / n_head
        opts["entry_gap"] += (extra * 0.45) / max(n_entry, 1)
        opts["docx_entry_after"] += min(4.0, extra * 0.12)
    return opts
